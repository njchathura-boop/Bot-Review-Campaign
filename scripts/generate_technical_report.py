from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "Bot_Review_Campaign_Technical_Report_Professional.docx"
DIAGRAM_DIR = ROOT / "reports" / "diagrams"


def _font(size):
    for name in ("arial.ttf", "segoeui.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def _diagram(name, title, rows):
    """Create a small SmartArt-like flow diagram as a PNG for Word."""
    width = 1500
    row_height = 140
    image = Image.new("RGB", (width, 150 + len(rows) * row_height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(38)
    text_font = _font(25)
    draw.text((40, 25), title, fill="#1f4e79", font=title_font)
    for row_index, row in enumerate(rows):
        y = 115 + row_index * row_height
        box_width = (width - 80 - (len(row) - 1) * 25) // len(row)
        for index, label in enumerate(row):
            x = 40 + index * (box_width + 25)
            fill = ["#d9eaf7", "#e2f0d9", "#fff2cc", "#fce4d6"][index % 4]
            draw.rounded_rectangle((x, y, x + box_width, y + 82), radius=18, fill=fill, outline="#4472c4", width=3)
            words = label.split()
            lines, current = [], ""
            for word in words:
                candidate = f"{current} {word}".strip()
                if draw.textbbox((0, 0), candidate, font=text_font)[2] > box_width - 25:
                    lines.append(current)
                    current = word
                else:
                    current = candidate
            if current:
                lines.append(current)
            total_height = len(lines) * 28
            for line_no, line in enumerate(lines):
                bbox = draw.textbbox((0, 0), line, font=text_font)
                tx = x + (box_width - (bbox[2] - bbox[0])) / 2
                ty = y + (82 - total_height) / 2 + line_no * 28
                draw.text((tx, ty), line, fill="#17365d", font=text_font)
            if index < len(row) - 1:
                ax = x + box_width + 4
                ay = y + 41
                draw.line((ax, ay, ax + 17, ay), fill="#c0504d", width=4)
                draw.polygon([(ax + 17, ay), (ax + 8, ay - 7), (ax + 8, ay + 7)], fill="#c0504d")
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    path = DIAGRAM_DIR / name
    image.save(path, "PNG")
    return path


def make_diagrams():
    return {
        "architecture": _diagram(
            "architecture.png", "End-to-end Detectra architecture",
            [
                ["Raw labelled reviews", "Amazon behaviour", "Catalogue"],
                ["DVC + preprocessing", "Temporal features", "Leakage-safe splits"],
                ["Airflow", "Ray Tune", "MLflow"],
                ["Model bundles", "FastAPI + UI", "Kafka/Spark stream"],
                ["Prometheus -> Grafana", "Filebeat -> Elasticsearch -> Kibana"],
            ],
        ),
        "streaming.png": _diagram(
            "streaming.png", "Online event-time campaign pipeline",
            [
                ["Producer replay / live reviews"],
                ["Kafka: reviews.raw.v1"],
                ["Spark windows + watermark"],
                ["analysis-windows.v1"],
                ["Graph + hybrid scorer"],
                ["campaign-scores.v1 -> FastAPI UI"],
            ],
        ),
        "cicd.png": _diagram(
            "cicd.png", "CI/CD and deployment flow",
            [
                ["Git commit / release tag"],
                ["GitHub Actions tests + Ruff + smoke"],
                ["Docker BuildKit + Trivy"],
                ["GHCR immutable Git-SHA images"],
                ["Kustomize + kubectl / Argo CD"],
                ["Kubernetes probes + rollout"],
            ],
        ),
    }


DIAGRAMS = make_diagrams()
if (DIAGRAM_DIR / "professional_architecture.png").exists():
    DIAGRAMS["architecture"] = DIAGRAM_DIR / "professional_architecture.png"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True)
        shade(t.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph()
    return t


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.style = "No Spacing"
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(45, 45, 45)
    return p


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def paragraph(doc, text):
    doc.add_paragraph(text)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
styles["Heading 2"].font.color.rgb = RGBColor(47, 84, 150)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Bot Review Campaign Detection")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(31, 78, 121)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Technical Report and Demonstration Guide")
r.italic = True
r.font.size = Pt(13)
doc.add_paragraph()
paragraph(doc, "Course project report prepared from the repository README, technical documentation, deployment manifests, and runnable pipeline configuration.")
table(doc, ["Item", "Value"], [
    ("Project", "Bot Review Campaign Detection"),
    ("Primary interface", "FastAPI + browser review/campaign console"),
    ("Training models", "TF-IDF Logistic Regression baseline, DistilBERT review model, hybrid campaign model"),
    ("MLOps components", "Airflow, Ray Tune, MLflow, DVC, Docker, Kubernetes, Prometheus, Grafana"),
    ("Streaming components", "Kafka, Spark Structured Streaming, campaign scorer"),
])
paragraph(doc, "Responsible-use note: a risk score is evidence for human review, not proof that a person is a bot. Campaign restrictions require moderator confirmation.")

heading(doc, "1. Problem statement and objectives")
paragraph(doc, "Ecommerce platforms receive large volumes of reviews. Some activity may be ordinary customer feedback, while other activity may contain generated wording, coordinated accounts, unusual timing, rating attacks, or promotional campaigns. A single text classifier cannot reliably identify coordination, and a group detector should not accuse a user based on one sentence alone.")
paragraph(doc, "This project therefore separates two decisions. Review risk estimates whether one review should be sent to moderator attention. Campaign risk evaluates a group of related reviews using language, account, product, rating, timing, and graph evidence. The system presents evidence and lineage to a moderator; it does not automatically punish an account.")
table(doc, ["Objective", "Implementation"], [
    ("Individual review triage", "Calibrated DistilBERT probability and evidence in FastAPI/UI."),
    ("Coordinated campaign detection", "Hybrid DistilBERT + numeric group model after Kafka/Spark event-time grouping."),
    ("Reproducible training", "DVC manifests, Ray Tune trials, MLflow lineage, held-out evaluation."),
    ("Operational delivery", "Docker Compose locally; Kubernetes/Argo CD manifests for staging."),
])

heading(doc, "2. System architecture")
paragraph(doc, "The architecture has an offline lane and an online lane. The offline lane prepares data and trains candidate models. The online lane accepts review scans and processes replayed review events. Monitoring and logging observe both lanes without changing model decisions.")
pic = doc.add_picture(str(DIAGRAMS["architecture"]), width=Inches(6.8))
pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
code(doc, "OFFLINE: raw data -> DVC/ETL -> Airflow -> Ray Tune -> MLflow -> approved model bundle")
code(doc, "ONLINE: browser -> FastAPI -> review model -> evidence response")
code(doc, "STREAM: producer -> Kafka raw topic -> Spark event-time windows -> campaign scorer -> Kafka scores -> FastAPI UI")
code(doc, "OBSERVE: API/Ray metrics -> Prometheus -> Grafana; Docker logs -> Filebeat -> Elasticsearch -> Kibana")
table(doc, ["Layer", "Main components", "Responsibility"], [
    ("Data", "Amazon Reviews 2023, labelled ecommerce reviews, catalogue", "Observed events, text labels, product/launch context."),
    ("Preparation", "data.py, temporal.py, labeled_datasets.py, synthetic.py", "Canonical schemas, UTC timestamps, past-only features, leakage-safe splits."),
    ("Training", "Airflow, Ray Tune, DistilBERT, hybrid MLP head, MLflow", "Schedule jobs, search hyperparameters, evaluate and register candidates."),
    ("Serving", "FastAPI, review/campaign bundles, browser UI", "Score reviews, expose evidence, accept moderator decisions."),
    ("Streaming", "Kafka, Spark, campaign_scorer.py", "Buffer events, group by event time, score coordinated windows."),
    ("Operations", "Docker, Kubernetes, Prometheus, Grafana, Elastic/Kibana", "Package, deploy, measure, chart, and investigate behavior."),
])
heading(doc, "2.1 System architecture diagram", 2)
code(doc, "                         OFFLINE TRAINING")
code(doc, "Amazon + labelled data -> DVC/ETL -> Airflow -> Ray Tune -> MLflow -> model bundles")
code(doc, "                                                           |                 |")
code(doc, "                                                           +--> review model + campaign model")
code(doc, "                                                                                  |")
code(doc, "BROWSER -> FastAPI ------------------------------------------------------------+")
code(doc, "             |                  |                          |")
code(doc, "             v                  v                          v")
code(doc, "       Review DistilBERT   Kafka raw -> Spark -> scorer -> Kafka scores -> API UI")
code(doc, "             |                                             ")
code(doc, "       Prometheus -> Grafana       Docker logs -> Filebeat -> Elasticsearch -> Kibana")

heading(doc, "3. Dataset description and preparation")
paragraph(doc, "The text and behavioral roles are deliberately separate. Labelled ecommerce review files provide text supervision. Amazon Reviews 2023 provides observed product, account, rating, timestamp, verification, category, and helpful-vote behavior, but does not provide trustworthy bot-campaign labels. The project does not invent labels for those real Amazon rows.")
paragraph(doc, "The loader maps source aliases into a canonical schema, normalizes text and IDs, converts timestamps to UTC, rejects malformed rows and duplicate IDs, and records provenance. Temporal features use only information available before each event: time since a prior review, previous product/user counts, hour, weekday, launch phase, and inter-arrival time. This prevents future leakage.")
paragraph(doc, "Controlled campaign scenarios are generated from distributions learned from observed Amazon behavior. They include organic activity, legitimate launch bursts, coordinated positive/negative campaigns, paraphrased campaigns, off-hour campaigns, slow-drip campaigns, and multi-product campaigns. Scenario groups stay intact within train, validation, or test splits.")
table(doc, ["Output", "Purpose", "Documented size/status"], [
    ("Text training", "Fine-tune review DistilBERT", "49,946 rows"),
    ("Real validation", "Calibrate threshold and select trial", "9,260 rows"),
    ("Real test", "Final individual-review evaluation", "9,168 rows"),
    ("Amazon behavioral source", "Past-only behavior profile", "816,216 downloaded rows in documented run"),
    ("Campaign splits", "Group-level hybrid model", "Manifest-controlled train/validation/test groups"),
])
paragraph(doc, "Synthetic text augmentation, when enabled, is train-only and capped by policy. Every generated bundle writes a manifest with schema version, seed, source provenance, row counts, and SHA-256 digests. DVC records the reproduction graph and lockfile; Git stores the code and metadata, not large raw datasets.")

heading(doc, "4. Data pipeline and streaming execution")
paragraph(doc, "Airflow is the timetable manager for expensive offline work. The retraining DAG submits temporal ETL to Ray, waits with a rescheduling sensor, submits review training, waits, then submits campaign training. It is manual by default because full tuning is expensive and should only run when enough approved data exists.")
paragraph(doc, "Kafka is the durable conveyor belt. The producer reads a controlled JSONL replay and publishes events to reviews.raw.v1. Kafka provides buffering, replay, back-pressure, and recovery. Spark Structured Streaming reads the raw topic, parses event time, applies windows and watermarks, validates fields, and publishes reviews.analysis-windows.v1. The campaign scorer consumes those windows, builds the related-review graph, invokes the hybrid model, and publishes reviews.campaign-scores.v1. The FastAPI-side consumer materializes those scores for the UI.")
pic = doc.add_picture(str(DIAGRAMS["streaming.png"]), width=Inches(6.8))
pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
code(doc, "docker compose --profile stream up -d spark-stream campaign-scorer")
code(doc, "docker compose exec -T campaign-scorer python streaming/producer.py --input /opt/project/data/processed/temporal_bundle/campaign/test.jsonl --bootstrap-servers kafka:29092 --rate 10")
code(doc, "docker compose exec kafka /opt/kafka/bin/kafka-console-consumer.sh --bootstrap-server kafka:29092 --topic reviews.campaign-scores.v1 --from-beginning")

heading(doc, "5. Model development")
heading(doc, "5.1 Baseline model", 2)
paragraph(doc, "The retained baseline is class-balanced Logistic Regression over TF-IDF word unigrams/bigrams and interpretable style indicators such as punctuation density, uppercase ratio, repetition, and promotional wording. It is fast, deterministic, and useful as a comparison and fallback. It is not the promoted UI model when the DistilBERT bundle is available.")
heading(doc, "5.2 Individual-review DistilBERT", 2)
paragraph(doc, "The promoted review model fine-tunes distilbert-base-uncased for two labels. DistilBERT was chosen because it captures meaning and context better than keywords while being smaller and faster than full BERT. Validation data calibrates the probability temperature and decision threshold; the test split is read only after trial selection. Ray Tune searches learning rate, weight decay, dropout, batch size, sequence length, frozen layers, warm-up, gradient clipping, class-weight multiplier, label smoothing, and epochs. MLflow records each trial and the selected bundle.")
heading(doc, "5.3 Hybrid campaign model", 2)
paragraph(doc, "The campaign model is a hybrid neural network. It encodes up to six review texts using DistilBERT, pools the text representations, passes numeric group features through a small multilayer perceptron, concatenates the two representations, and sends them through fusion layers to a sigmoid campaign-risk output. Numeric features include review/user/product counts, duration, review rate, inter-arrival statistics, rating concentration, verified ratio, helpful votes, off-hour/weekend activity, and launch proximity. Scenario IDs, future-derived values, and target labels are excluded from model input.")
paragraph(doc, "This structure was selected because language alone misses coordination and raw behavior numbers cannot understand wording. The fusion head lets each source contribute while preserving a clear evidence trail. Related reviews are connected by shared product/user context, close event-time windows, and semantic similarity edges; connected components become candidate groups.")
table(doc, ["Evaluation measure", "Why it matters"], [
    ("PR-AUC", "Useful for imbalanced suspicious-review/campaign decisions."),
    ("ROC-AUC", "General ranking quality across thresholds."),
    ("Precision/Recall/F1", "Trade-off between moderator workload and missed suspicious activity."),
    ("Brier score / log loss", "Probability quality and calibration."),
    ("Legitimate-burst false-alert rate", "Checks that busy but genuine launches are not over-flagged."),
])

heading(doc, "6. Experiment tracking, versioning, and orchestration")
paragraph(doc, "Ray Tune distributes hyperparameter trials and can resume checkpoints. MLflow is the experiment diary and model-lineage registry: it stores trial parameters, validation metrics, artifacts, Git SHA, data hashes, and the selected model identity. Airflow controls when the sequence runs but does not replace Ray or MLflow. DVC stores data-stage dependencies, file hashes, and reproducible commands in dvc.yaml/dvc.lock. Git tracks source, tests, manifests, and deployment files; Git LFS stores large promoted model binaries.")
code(doc, "docker compose -f orchestration/docker-compose.airflow.yml up airflow-init")
code(doc, "docker compose -f orchestration/docker-compose.airflow.yml up -d airflow-dag-processor airflow-scheduler airflow-api-server")
code(doc, "docker compose -f orchestration/docker-compose.airflow.yml exec airflow-api-server airflow dags trigger bot_campaign_model_retraining")
paragraph(doc, "Airflow is available at http://localhost:8084 with the local development login admin/admin. MLflow is at http://localhost:5001 and Ray at http://localhost:8265.")

heading(doc, "7. API, Docker, and Kubernetes deployment")
paragraph(doc, "FastAPI serves the static UI and versioned endpoints. The main endpoints are POST /v1/reviews/score, POST /v1/reviews/batch-score, POST /v1/demo/replay, GET /v1/campaigns, POST /v1/campaigns/{id}/decision, GET /v1/ops/summary, GET /v1/lineage/current, GET /health/ready, and GET /metrics. The UI uses the API; it does not directly access Kafka, Spark, Ray, or model files.")
paragraph(doc, "Docker Compose is the laptop runtime. Dockerfile builds the API/UI image and docker/ray.Dockerfile builds the jobs/training image. Compose supplies internal service names, ports, health checks, volumes, and optional profiles. The GPU overlay exposes the single laptop GPU to ray-worker only; the normal stack remains CPU-portable.")
paragraph(doc, "Kubernetes is the cluster runtime. k8s/ contains Deployments, Services, CronJobs, ConfigMaps, persistent claims, monitoring, and GPU overlays. Probes prevent traffic from reaching an unready API. Resource limits and non-root security settings make cluster execution safer. Argo CD can synchronize the manifests from Git.")
table(doc, ["Image repository", "Use"], [
    ("ghcr.io/njchathura-boop/bot-review-campaign-api", "FastAPI/UI runtime image"),
    ("ghcr.io/njchathura-boop/bot-review-campaign-jobs", "Ray ETL/training image"),
])
paragraph(doc, "GitHub Actions builds and Trivy-scans both images, pushes a full Git-SHA tag plus a convenience latest tag to GitHub Container Registry, renders Kustomize with the immutable SHA tag, and deploys staging with kubectl. Source code and manifests are committed to Git; built images are pushed to GHCR rather than stored inside Git.")

heading(doc, "8. Monitoring and logging")
paragraph(doc, "Prometheus scrapes API and Ray /metrics endpoints every 10 seconds. Grafana turns those time-series into charts for health, scan rate, p50/p95/p99 latency, API errors, model risk/confidence, campaign workload, moderator decisions, Ray CPU/memory, and streaming activity. These metrics are operational signals, not model decisions.")
paragraph(doc, "Filebeat reads Docker JSON logs, adds container metadata, and sends them to Elasticsearch. Kibana searches the resulting filebeat-* data view using KQL. Grafana answers “is the system slow?” while Kibana answers “what exact log explains it?” Example KQL is container.name : \"*spark*\" AND message : \"ERROR\".")
table(doc, ["Metric family", "Example interpretation"], [
    ("API latency", "High p95 means some users are waiting too long even if average latency looks fine."),
    ("API errors", "A rise in 4xx/5xx indicates invalid clients or service failures."),
    ("Review model", "Risk/confidence distributions show how the model behaves over time."),
    ("Campaign workload", "Candidate and human-review counts show moderator load."),
    ("Ray health", "CPU/memory/GPU trends reveal resource pressure during training."),
])

heading(doc, "9. CI/CD and reproducibility")
paragraph(doc, "The CI workflow runs installation, linting, tests, UI/API contracts, training smoke checks, Docker BuildKit packaging, and vulnerability scanning. Release tags trigger CD. CD checks out Git LFS model bundles, builds API and jobs images, scans them with Trivy, pushes immutable images to GHCR, renders Kubernetes manifests, waits for rollouts, and calls /health/ready plus /metrics as a smoke test. DVC repro is explicit: changing a data input does not silently retrain a model; an operator or scheduled workflow must run dvc repro and then approve a new training run.")
pic = doc.add_picture(str(DIAGRAMS["cicd.png"]), width=Inches(6.8))
pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
code(doc, "dvc repro")
code(doc, "git tag v1.0.0")
code(doc, "git push origin v1.0.0")

heading(doc, "10. Results, challenges, and limitations")
paragraph(doc, "The documented promoted review trial achieved PR-AUC 0.9353, ROC-AUC 0.9215, F1 0.8308, precision 0.8005, recall 0.8635, Brier score 0.1305, and log loss 0.4443 on its recorded validation evaluation. The report distinguishes this from a final untouched test result and from controlled campaign-scenario scores. A tiny four-row baseline smoke artifact must not be presented as a statistically meaningful comparison.")
paragraph(doc, "The main challenges were large resumable downloads, temporal leakage, group leakage, laptop memory pressure during Ray/DistilBERT training, Windows execution policy, model cold-start time, and confusion between internal Docker names and browser localhost names. Solutions included bounded downloads, past-only features, complete-group splits, four-gigabyte Ray shared memory, checkpointing, lazy model loading, explicit port mappings, and health/lineage endpoints.")
paragraph(doc, "The most important limitation is ground truth: public behavioral data does not contain verified campaign labels. Controlled scenarios validate the coordination pipeline but do not prove real-world abuse prevalence. Future work should collect moderator-confirmed labels, add multilingual models, secure model artifact provisioning, use a shared online feature store, add privacy controls, run shadow-mode canaries, and evaluate moderator workload before enforcement.")

heading(doc, "11. Demonstration checklist")
for item in [
    "Show the README and repository folders, then explain the architecture diagram.",
    "Start Docker Compose and show docker compose ps plus the FastAPI UI/API docs.",
    "Trigger the Airflow retraining DAG and show the ordered tasks.",
    "Show Ray trials/resources and the corresponding MLflow experiment/model lineage.",
    "Run the Kafka replay and show Spark event-time processing and campaign scores.",
    "Show a review scan and the human-review evidence warning in the UI.",
    "Open Grafana for latency/model/campaign metrics and Kibana for a correlated log using KQL.",
    "Show the GitHub Actions image scan and the immutable GHCR Git-SHA image reference.",
]:
    bullet(doc, item)
paragraph(doc, "This sequence demonstrates the complete required scope: Airflow automation, Kafka/Spark data engineering, preprocessing, multiple models, MLflow, DVC/versioning, Docker deployment, monitoring/logging, Kubernetes manifests, and CI/CD.")

doc.add_page_break()
heading(doc, "Appendix A. Key repository structure")
table(doc, ["Path", "Purpose"], [
    ("src/bot_campaign/", "Canonical schemas, data loading, features, models, runtime, API routes, observability."),
    ("training/", "Ray Tune review and campaign training programs."),
    ("streaming/", "Kafka producer and campaign scorer."),
    ("spark/", "Spark Structured Streaming review-window job."),
    ("orchestration/", "Airflow Compose file, password file, DAGs."),
    ("monitoring/", "Prometheus configuration, Grafana dashboards, Filebeat configuration."),
    ("data/", "Raw/processed data contracts and beginner dataset guide."),
    ("artifacts/", "Model bundles and candidate outputs; large promoted binaries use Git LFS."),
    ("dvc.yaml / dvc.lock", "Data-stage reproduction graph and exact dependency hashes."),
    ("Dockerfile / docker/ray.Dockerfile", "API and training image definitions."),
    ("docker-compose.yml", "Local multi-service runtime."),
    ("k8s/ and deploy/", "Kustomize cluster manifests and Argo CD application."),
    (".github/workflows/", "CI quality gates and CD image/deployment workflow."),
])
paragraph(doc, "End of report.")

doc.core_properties.title = "Bot Review Campaign Detection - Technical Report"
doc.core_properties.subject = "MLOps technical report"
doc.core_properties.author = "Bot Review Campaign Project"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
